"""入库主流程：加载 → 切分 → 向量化 → 写入 Qdrant。"""

import json
import os
from pathlib import Path
from typing import Any

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from src.core.logger import get_logger
from src.rag.config import get_rag_settings
from src.rag.vector_store import content_hash, ensure_collection, upsert_chunks

logger = get_logger(__name__)

# 本地语料目录中扫描的文件后缀：JSONL（每行一条）+ JSON 数组（整体一个文件）
CORPUS_FILE_TYPES: tuple[str, ...] = (".jsonl", ".json")
# 在线附件：txt / pdf / docx
ATTACHMENT_TYPES: tuple[str, ...] = (".txt", ".pdf", ".docx")

# 从 JSON 记录中提取正文时依次尝试的字段名
# 已知数据集：data/raws/law_library.jsonl → content；data/cases/*.json → fact
_TEXT_KEYS: tuple[str, ...] = ("text", "content", "fact", "正文", "内容", "sentence")


def list_files(path: str | Path, allowed_types: tuple[str, ...]) -> tuple[str, ...]:
    """列出文件夹内后缀匹配的文件；路径非法或无匹配时返回空元组。"""
    if not os.path.isdir(path):
        logger.error(f"[list_files]{path} 不是文件夹")
        return ()

    return tuple(
        os.path.join(path, name)
        for name in sorted(os.listdir(path))
        if name.endswith(allowed_types)
    )


def _to_document(
    record: Any, source: str, index: int, text_key: str | None = None
) -> Document | None:
    """把一条 JSON 记录转成 Document；无法提取正文时返回 None。"""
    if isinstance(record, str):
        content, extra = record, {}
    elif isinstance(record, dict):
        keys = (text_key,) if text_key else _TEXT_KEYS

        matched: str | None = None
        content = ""
        for key in keys:
            value = record.get(key)
            if isinstance(value, str) and value.strip():
                matched = key
                content = value
                break

        # 正文字段已提取，其余字段原样保留为 metadata
        extra = {k: v for k, v in record.items() if k != matched}
    else:
        logger.warning(
            f"[_to_document]{source} 第 {index} 条记录类型不支持：{type(record).__name__}"
        )
        return None

    if not content.strip():
        logger.warning(f"[_to_document]{source} 第 {index} 条记录未找到正文字段，已跳过")
        return None

    return Document(page_content=content, metadata={**extra, "source": source, "index": index})


def _records_to_documents(
    records: list, source: str, text_key: str | None = None
) -> list[Document]:
    """批量把记录转成 Document，跳过无效记录。"""
    documents = []
    for index, record in enumerate(records):
        document = _to_document(record, source, index, text_key)
        if document is not None:
            documents.append(document)
    return documents


def load_jsonl(filepath: str | Path, text_key: str | None = None) -> list[Document]:
    """加载 JSONL 文件（每行一条记录）；单行解析失败只跳过该行。"""
    filepath = Path(filepath)
    records = []

    with open(filepath, "r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue

            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                logger.warning(f"[load_jsonl]{filepath.name} 第 {line_no} 行解析失败，已跳过：{e}")

    documents = _records_to_documents(records, filepath.name, text_key)
    logger.info(f"[load_jsonl]{filepath.name} 加载 {len(documents)} 条")
    return documents


def _is_jsonl(filepath: Path) -> bool:
    """
    探测文件是否为 JSONL。

    首个非空行能单独解析成 JSON 对象即为 JSONL；否则视为整体一个 JSON（数组或大对象）。
    有些数据集会把 JSONL 存成 .json（如 CAIL 裁判文书），所以不能只看后缀。
    """
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue

            try:
                return isinstance(json.loads(line), dict)
            except json.JSONDecodeError:
                return False

    return False


def load_json(filepath: str | Path, text_key: str | None = None) -> list[Document]:
    """加载 JSON 文件（整体一个数组或单个对象）。"""
    filepath = Path(filepath)

    if _is_jsonl(filepath):
        logger.debug(f"[load_json]{filepath.name} 内容实为 JSONL，转按行解析")
        return load_jsonl(filepath, text_key)

    try:
        with open(filepath, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        logger.error(f"[load_json]{filepath.name} 解析失败：{e}")
        raise

    records = data if isinstance(data, list) else [data]
    documents = _records_to_documents(records, filepath.name, text_key)
    logger.info(f"[load_json]{filepath.name} 加载 {len(documents)} 条")
    return documents


def load_txt(filepath: str | Path) -> list[Document]:
    """加载 TXT 文件。"""
    return TextLoader(str(filepath), encoding="utf-8").load()


def load_pdf(filepath: str | Path, password: str | None = None) -> list[Document]:
    """加载 PDF 文件。"""
    return PyPDFLoader(str(filepath), password=password).load()


def load_docx(filepath: str | Path) -> list[Document]:
    """加载 DOCX 文件。"""
    try:
        from docx import Document as DocxDocument
    except ImportError as e:
        logger.error("[load_docx]缺少 python-docx，请先安装依赖")
        raise e

    filepath = Path(filepath)
    content = "\n".join(p.text for p in DocxDocument(str(filepath)).paragraphs)

    if not content.strip():
        logger.warning(f"[load_docx]{filepath.name} 未解析出文本内容")
        return []

    return [Document(page_content=content, metadata={"source": filepath.name})]


def load_file(filepath: str | Path, text_key: str | None = None) -> list[Document]:
    """按后缀分发到对应加载器，text_key 仅对 JSON/JSONL 生效。"""
    filepath = Path(filepath)
    suffix = filepath.suffix.lower()

    if suffix == ".jsonl":
        return load_jsonl(filepath, text_key)
    if suffix == ".json":
        return load_json(filepath, text_key)
    if suffix == ".txt":
        return load_txt(filepath)
    if suffix == ".pdf":
        return load_pdf(filepath)
    if suffix == ".docx":
        return load_docx(filepath)
    if suffix == ".doc":
        raise ValueError(f"不支持旧版 .doc：{filepath.name}，请先另存为 .docx")

    raise ValueError(f"不支持的文件类型：{suffix}（{filepath.name}）")


def load_dir(
    path: str | Path,
    allowed_types: tuple[str, ...] = CORPUS_FILE_TYPES,
    text_key: str | None = None,
) -> list[Document]:
    """加载文件夹内所有后缀匹配的文件。"""
    documents = []
    for filepath in list_files(path, allowed_types):
        documents.extend(load_file(filepath, text_key))

    logger.info(f"[load_dir]{path} 共加载 {len(documents)} 条")
    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    """按 configs/rag.yaml 的 chunking 配置切分文档。"""
    settings = get_rag_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunking.chunk_size,
        chunk_overlap=settings.chunking.chunk_overlap,
    )
    return splitter.split_documents(documents)


def dedup_documents(documents: list[Document]) -> list[Document]:
    """
    按内容哈希精确去重，保留首次出现的记录。

    仅做精确比对，不使用相似度去重：法律文本一字之差即不同罪名与刑期，
    近似去重会误删真实法条（如各部法律中同款的"构成犯罪的，依法追究
    刑事责任"）与不同被告人的案件样本。
    """
    seen: set[str] = set()
    kept: list[Document] = []

    for document in documents:
        digest = content_hash(document)
        if digest in seen:
            continue
        seen.add(digest)
        kept.append(document)

    removed = len(documents) - len(kept)
    if removed:
        logger.info(f"[dedup_documents] 去重剔除 {removed} 条，剩余 {len(kept)} 条")

    return kept


def _collection_of(corpus: str) -> str:
    """把语料类型映射为 Qdrant collection 名。"""
    return get_rag_settings().collection_of(corpus)


def ingest(
    path: str | Path,
    corpus: str,
    text_key: str | None = None,
    limit: int | None = None,
) -> int:
    """
    入库主流程：加载 → 去重 → 切分 → 向量化 → 写入 Qdrant。

    以内容哈希作为 point id，同内容重跑会覆盖同一个点，因此可安全重试。
    """
    path = Path(path)

    if path.is_file():
        documents = load_file(path, text_key)
    else:
        documents = load_dir(path, CORPUS_FILE_TYPES, text_key)

    documents = dedup_documents(documents)
    chunks = split_documents(documents)

    if limit is not None:
        chunks = chunks[:limit]

    collection = _collection_of(corpus)
    logger.info(f"[ingest]{path} 加载 {len(documents)} 条，切分 {len(chunks)} 块 → {collection}")

    ensure_collection(collection)

    return upsert_chunks(collection, chunks)


if __name__ == "__main__":
    from src.core.paths import CASES_DIR, RAWS_DIR

    # 小样本试跑：各 1000 块，确认限流、维度与写入均正常后再全量入库
    ingest(RAWS_DIR, "laws", limit=1000)
    ingest(CASES_DIR / "data_train.json", "cases", limit=1000)